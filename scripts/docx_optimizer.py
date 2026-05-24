#!/usr/bin/env python3
"""
简历优化主程序
功能：读取原始 DOCX，应用优化建议，添加详细批注
"""

import sys
import json
import argparse
from pathlib import Path
from docx import Document
from docx.shared import RGBColor, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from datetime import datetime
import shutil


class ResumeOptimizer:
    """简历优化器"""

    def __init__(self, original_path, output_dir):
        self.original_path = Path(original_path)
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        # 加载文档
        self.doc = Document(original_path)
        self.comment_id = 0

        # 批注存储
        self.comments = []

    def apply_optimizations(self, optimizations):
        """
        应用优化建议

        Args:
            optimizations: list of dict, 每个包含：
                - paragraph_index: 段落索引
                - run_index: run 索引（可选）
                - old_text: 原文
                - new_text: 新文本
                - comment: 批注内容
                - modification_type: 修改类型（关键词优化、量化成果等）
        """
        for opt in optimizations:
            try:
                self._apply_single_optimization(opt)
            except Exception as e:
                print(f"⚠️  应用优化失败: {opt.get('old_text', '')[:30]}... - {e}")
                continue

    def _apply_single_optimization(self, opt):
        """应用单个优化"""
        para_idx = opt.get('paragraph_index')
        old_text = opt.get('old_text', '')
        new_text = opt.get('new_text', '')
        comment_text = opt.get('comment', '')

        # 查找段落
        if para_idx is not None and para_idx < len(self.doc.paragraphs):
            para = self.doc.paragraphs[para_idx]
        else:
            # 如果没有指定段落，尝试全文搜索
            para = self._find_paragraph_by_text(old_text)
            if not para:
                print(f"⚠️  未找到段落: {old_text[:50]}...")
                return

        # 替换文本
        original_text = para.text
        if old_text in original_text:
            # 保留原有格式，替换文本
            new_para_text = original_text.replace(old_text, new_text)

            # 清空段落内容但保留格式
            for run in para.runs:
                run.text = ''

            # 添加新文本（使用第一个 run 的格式）
            if para.runs:
                para.runs[0].text = new_para_text
            else:
                para.add_run(new_para_text)

            # 添加批注
            if comment_text:
                self._add_comment_to_paragraph(para, comment_text)
        else:
            print(f"⚠️  文本不匹配: {old_text[:50]}...")

    def _find_paragraph_by_text(self, text):
        """根据文本内容查找段落"""
        for para in self.doc.paragraphs:
            if text in para.text:
                return para
        return None

    def _add_comment_to_paragraph(self, paragraph, comment_text):
        """
        在段落末尾添加批注

        注意：python-docx 不直接支持批注，这里使用一个变通方案：
        在段落末尾添加一个带特殊格式的 run 来标识批注内容
        """
        # 增加批注 ID
        self.comment_id += 1

        # 在段落末尾添加批注标记
        comment_mark = paragraph.add_run(f" [批注 {self.comment_id}]")
        comment_mark.font.color.rgb = RGBColor(255, 0, 0)  # 红色
        comment_mark.font.size = Pt(9)
        comment_mark.font.italic = True

        # 存储批注内容（稍后添加到文档末尾）
        self.comments.append({
            'id': self.comment_id,
            'text': comment_text,
            'paragraph': paragraph
        })

    def _add_comments_section(self):
        """在文档末尾添加批注汇总区域"""
        if not self.comments:
            return

        # 添加分页符
        self.doc.add_page_break()

        # 添加批注汇总标题
        title = self.doc.add_paragraph()
        title_run = title.add_run('📝 优化说明汇总')
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_run.font.color.rgb = RGBColor(0, 102, 204)

        # 添加说明
        note = self.doc.add_paragraph()
        note.add_run('以下是所有优化位置的详细说明。查看完毕后，可删除此页。').font.size = Pt(10)

        # 添加分隔线
        self.doc.add_paragraph('━' * 50)

        # 添加每个批注
        for comment in self.comments:
            # 批注 ID
            id_para = self.doc.add_paragraph()
            id_run = id_para.add_run(f"[批注 {comment['id']}]")
            id_run.bold = True
            id_run.font.color.rgb = RGBColor(255, 0, 0)

            # 批注内容
            content_para = self.doc.add_paragraph()
            content_para.add_run(comment['text']).font.size = Pt(10)

            # 分隔线
            self.doc.add_paragraph('─' * 40)

    def save(self, job_title="职位"):
        """保存优化后的文档"""
        # 添加批注汇总区域
        self._add_comments_section()

        # 保存原始备份
        backup_path = self.output_dir / 'resume_original_backup.docx'
        shutil.copy2(self.original_path, backup_path)

        # 保存优化版本
        optimized_path = self.output_dir / f'resume_optimized_{job_title}.docx'
        self.doc.save(optimized_path)

        return {
            'backup': str(backup_path),
            'optimized': str(optimized_path),
            'comment_count': len(self.comments)
        }


def main():
    """主函数"""
    parser = argparse.ArgumentParser(description='简历优化工具')
    parser.add_argument('--original', required=True, help='原始简历路径')
    parser.add_argument('--optimizations', required=True, help='优化建议 JSON 文件')
    parser.add_argument('--output-dir', required=True, help='输出目录')
    parser.add_argument('--job-title', default='职位', help='目标职位名称')

    args = parser.parse_args()

    # 加载优化建议
    with open(args.optimizations, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # 兼容多种 JSON 格式
    if isinstance(data, list):
        optimizations = data
    elif isinstance(data, dict):
        if 'optimizations' in data:
            optimizations = data['optimizations']
        else:
            optimizations = None
            for v in data.values():
                if isinstance(v, list):
                    optimizations = v
                    break
            if optimizations is None:
                print("❌ JSON 格式错误：找不到优化建议列表")
                sys.exit(1)
    else:
        print("❌ JSON 格式错误：期望列表或对象")
        sys.exit(1)

    # 验证每个优化项是字典
    optimizations = [opt for opt in optimizations if isinstance(opt, dict)]
    if not optimizations:
        print("❌ 没有有效的优化建议")
        sys.exit(1)

    # 创建优化器
    optimizer = ResumeOptimizer(args.original, args.output_dir)

    # 应用优化
    print("🔧 正在应用优化建议...")
    optimizer.apply_optimizations(optimizations)

    # 保存文档
    print("💾 正在保存文档...")
    result = optimizer.save(args.job_title)

    # 输出结果
    print("\n✅ 优化完成！")
    print(f"├─ 原始备份: {result['backup']}")
    print(f"├─ 优化版本: {result['optimized']}")
    print(f"└─ 批注数量: {result['comment_count']}")

    # 输出 JSON 结果供 Claude 读取
    print("\n__RESULT_JSON__")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
