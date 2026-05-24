#!/usr/bin/env python3
"""
简历优化高级版 - 支持真正的 Word 批注（Comments）
直接操作 OOXML 在修改位置添加 Word 原生批注
"""

import sys
import json
import argparse
from pathlib import Path
from docx import Document
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from datetime import datetime
import shutil
from lxml import etree
from copy import deepcopy


# Word comments 的 XML 命名空间
COMMENTS_URI = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments'
WML_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
WML_NS_MAP = {'w': WML_NS}


class CommentsPart:
    """管理 comments.xml 部分"""

    def __init__(self, doc):
        self.doc = doc
        self.comments_element = None
        self._init_comments()

    def _init_comments(self):
        """初始化或获取 comments.xml"""
        # 尝试获取已有的 comments part
        doc_part = self.doc.part
        try:
            for rel in doc_part.rels.values():
                if rel.reltype == COMMENTS_URI:
                    # 已有 comments part，解析它
                    self.comments_element = etree.fromstring(rel.target_part.blob)
                    self._comments_part = rel.target_part
                    return
        except Exception:
            pass

        # 创建新的 comments.xml
        comments_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:comments xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
            'xmlns:cx="http://schemas.microsoft.com/office/drawing/2014/chartex" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'xmlns:o="urn:schemas-microsoft-com:office:office" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
            'xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:w10="urn:schemas-microsoft-com:office:word" '
            'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
            'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" '
            'xmlns:w16cex="http://schemas.microsoft.com/office/word/2018/wordml/cex" '
            'xmlns:w16cid="http://schemas.microsoft.com/office/word/2016/wordml/cid" '
            'xmlns:w16="http://schemas.microsoft.com/office/word/2018/wordml" '
            'xmlns:w16se="http://schemas.microsoft.com/office/word/2015/wordml/symex" '
            'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
            'xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk" '
            'xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" '
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            '</w:comments>'
        )
        self.comments_element = etree.fromstring(comments_xml.encode('utf-8'))

        # 创建 Part 并添加关系
        comments_part = Part(
            partname=PackURI('/word/comments.xml'),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml',
            blob=etree.tostring(self.comments_element, xml_declaration=True, encoding='UTF-8', standalone=True),
            package=doc_part.package
        )
        doc_part.relate_to(comments_part, COMMENTS_URI)
        self._comments_part = comments_part

    def add_comment(self, comment_id, author, date_str, text):
        """添加一条批注到 comments.xml"""
        # 创建 <w:comment> 元素
        comment_elem = OxmlElement('w:comment')
        comment_elem.set(qn('w:id'), str(comment_id))
        comment_elem.set(qn('w:author'), author)
        comment_elem.set(qn('w:date'), date_str)
        comment_elem.set(qn('w:initials'), author[0] if author else 'A')

        # 将批注文本按行分段
        lines = text.split('\n')
        for line in lines:
            p_elem = OxmlElement('w:p')
            r_elem = OxmlElement('w:r')
            t_elem = OxmlElement('w:t')
            t_elem.set(qn('xml:space'), 'preserve')
            t_elem.text = line
            r_elem.append(t_elem)
            p_elem.append(r_elem)
            comment_elem.append(p_elem)

        self.comments_element.append(comment_elem)

    def save(self):
        """保存 comments.xml"""
        self._comments_part._blob = etree.tostring(
            self.comments_element,
            xml_declaration=True,
            encoding='UTF-8',
            standalone=True
        )


def add_comment_to_paragraph(paragraph, comment_id):
    """在段落中插入 commentRangeStart、commentRangeEnd 和 commentReference"""
    p_elem = paragraph._element

    # commentRangeStart — 插入到段落开头
    range_start = OxmlElement('w:commentRangeStart')
    range_start.set(qn('w:id'), str(comment_id))
    p_elem.insert(0, range_start)

    # commentRangeEnd — 插入到段落末尾
    range_end = OxmlElement('w:commentRangeEnd')
    range_end.set(qn('w:id'), str(comment_id))
    p_elem.append(range_end)

    # commentReference run — 添加批注引用标记
    ref_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rStyle = OxmlElement('w:rStyle')
    rStyle.set(qn('w:val'), 'CommentReference')
    rPr.append(rStyle)
    ref_run.append(rPr)
    ref_elem = OxmlElement('w:commentReference')
    ref_elem.set(qn('w:id'), str(comment_id))
    ref_run.append(ref_elem)
    p_elem.append(ref_run)


class ResumeOptimizerAdvanced:
    """简历优化器（高级版 - 真正的 Word 批注）"""

    def __init__(self, original_path, output_dir):
        self.original_path = Path(original_path)
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        # 加载文档
        self.doc = Document(original_path)

        # 初始化 comments part
        self.comments_part = CommentsPart(self.doc)

        # 批注计数
        self.comment_id = 0

        # 优化统计
        self.stats = {
            'total_optimizations': 0,
            'applied': 0,
            'skipped': 0,
            'keyword_optimizations': 0,
            'quantification': 0,
            'structure_improvements': 0,
            'wording_improvements': 0
        }

    def _collect_all_paragraphs(self):
        """收集文档中所有段落（包括表格中的段落）"""
        paragraphs = []
        # 正文段落
        for para in self.doc.paragraphs:
            paragraphs.append(para)
        # 表格中的段落
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        paragraphs.append(para)
        return paragraphs

    def apply_optimizations(self, optimizations):
        """应用优化建议"""
        self.stats['total_optimizations'] = len(optimizations)
        all_paragraphs = self._collect_all_paragraphs()
        now = datetime.now().strftime('%Y-%m-%dT%H:%M:%S+08:00')

        for opt in optimizations:
            try:
                old_text = opt.get('old_text', '')
                new_text = opt.get('new_text', '')
                comment_text = opt.get('comment', '')

                if not old_text:
                    self.stats['skipped'] += 1
                    continue

                # 查找匹配的段落
                matched_para = None
                for para in all_paragraphs:
                    if old_text in para.text:
                        matched_para = para
                        break

                if not matched_para:
                    # 尝试模糊匹配（去除空格差异）
                    old_stripped = old_text.replace(' ', '').replace('\u3000', '').replace('\t', '')
                    for para in all_paragraphs:
                        para_stripped = para.text.replace(' ', '').replace('\u3000', '').replace('\t', '')
                        if old_stripped in para_stripped:
                            matched_para = para
                            break

                if not matched_para:
                    print(f"⚠️  未找到匹配段落: {old_text[:50]}...")
                    self.stats['skipped'] += 1
                    continue

                # 替换文本（保留第一个 run 的格式）
                if new_text:  # 有替换文本
                    original_text = matched_para.text
                    new_para_text = original_text.replace(old_text, new_text)

                    if matched_para.runs:
                        # 保存第一个 run 的格式
                        first_run_format = matched_para.runs[0].font
                        # 清空所有 runs
                        for run in matched_para.runs:
                            run.text = ''
                        matched_para.runs[0].text = new_para_text
                    else:
                        matched_para.text = new_para_text
                else:
                    # new_text 为空表示删除（但保留段落，清空文本）
                    for run in matched_para.runs:
                        run.text = ''

                # 添加真正的 Word 批注
                if comment_text:
                    self.comment_id += 1
                    self.comments_part.add_comment(
                        comment_id=self.comment_id,
                        author='简历优化助手',
                        date_str=now,
                        text=comment_text
                    )
                    add_comment_to_paragraph(matched_para, self.comment_id)

                self.stats['applied'] += 1

                # 统计分类
                mod_type = opt.get('modification_type', '')
                if '关键词' in mod_type:
                    self.stats['keyword_optimizations'] += 1
                if '量化' in mod_type:
                    self.stats['quantification'] += 1
                if '结构' in mod_type:
                    self.stats['structure_improvements'] += 1
                if '措辞' in mod_type:
                    self.stats['wording_improvements'] += 1

            except Exception as e:
                print(f"⚠️  应用优化失败: {opt.get('old_text', '')[:30]}... - {e}")
                self.stats['skipped'] += 1

    def save(self, job_title="职位"):
        """保存文档"""
        # 保存 comments part
        self.comments_part.save()

        # 保存原始备份
        backup_path = self.output_dir / 'resume_original_backup.docx'
        shutil.copy2(self.original_path, backup_path)

        # 保存优化版本
        optimized_path = self.output_dir / f'resume_optimized_{job_title}.docx'
        self.doc.save(optimized_path)

        return {
            'backup': str(backup_path),
            'optimized': str(optimized_path),
            'stats': self.stats,
            'comment_count': self.comment_id
        }


def main():
    """主函数"""
    parser = argparse.ArgumentParser(description='简历优化工具（高级版 - Word 批注）')
    parser.add_argument('--original', required=True, help='原始简历路径')
    parser.add_argument('--optimizations', required=True, help='优化建议 JSON 文件')
    parser.add_argument('--output-dir', required=True, help='输出目录')
    parser.add_argument('--job-title', default='职位', help='目标职位名称')

    args = parser.parse_args()

    # 加载优化建议
    with open(args.optimizations, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # 兼容多种 JSON 格式
    if isinstance(data, list):
        optimizations = data
    elif isinstance(data, dict):
        if 'optimizations' in data:
            optimizations = data['optimizations']
        else:
            optimizations = None
            for v in data.values():
                if isinstance(v, list):
                    optimizations = v
                    break
            if optimizations is None:
                print("❌ JSON 格式错误：找不到优化建议列表")
                sys.exit(1)
    else:
        print("❌ JSON 格式错误：期望列表或对象")
        sys.exit(1)

    # 验证每个优化项是字典
    valid_optimizations = []
    for i, opt in enumerate(optimizations):
        if isinstance(opt, dict):
            valid_optimizations.append(opt)
        else:
            print(f"⚠️  跳过第 {i+1} 条优化：格式不正确（期望对象，实际为 {type(opt).__name__}）")
    optimizations = valid_optimizations

    if not optimizations:
        print("❌ 没有有效的优化建议")
        sys.exit(1)

    print(f"📋 加载了 {len(optimizations)} 条优化建议")

    # 创建优化器
    print("📂 加载简历文档...")
    optimizer = ResumeOptimizerAdvanced(args.original, args.output_dir)

    # 应用优化
    print(f"🔧 正在应用优化建议...")
    optimizer.apply_optimizations(optimizations)

    # 保存文档
    print("💾 正在保存文档...")
    result = optimizer.save(args.job_title)

    # 输出结果
    print("\n" + "=" * 60)
    print("✅ 简历优化完成！")
    print("=" * 60)
    stats = result['stats']
    print(f"\n📊 优化统计:")
    print(f"  ├─ 总优化数: {stats['total_optimizations']}")
    print(f"  ├─ 成功应用: {stats['applied']}")
    print(f"  ├─ 跳过: {stats['skipped']}")
    print(f"  ├─ 关键词优化: {stats['keyword_optimizations']}")
    print(f"  ├─ 量化成果: {stats['quantification']}")
    print(f"  ├─ 结构调整: {stats['structure_improvements']}")
    print(f"  └─ 措辞优化: {stats['wording_improvements']}")
    print(f"\n📁 输出文件:")
    print(f"  ├─ 原始备份: {result['backup']}")
    print(f"  └─ 优化版本: {result['optimized']} ⭐")
    print(f"\n💬 Word 批注数量: {result['comment_count']}")
    print("\n💡 用 Word 打开优化版本，批注会显示在右侧边栏")
    print("=" * 60)

    # 输出 JSON 结果
    print("\n__RESULT_JSON__")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
